# -*- coding: utf-8 -*-
"""Build 钢铁防线 3D · 坦克档案 (DOCX) from game data and captured renders."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
A = lambda p: os.path.join(ROOT, 'docs', 'design', 'assets', p)

NAVY = RGBColor(0x14, 0x1d, 0x2e)
NAVY_HEX = '141D2E'
GOLD = RGBColor(0xb0, 0x7d, 0x22)
GOLD_HEX = 'B07D22'
GOLD_LIGHT = RGBColor(0xc9, 0x9a, 0x3e)
BODY = RGBColor(0x2b, 0x30, 0x38)
GRAY = RGBColor(0x6b, 0x74, 0x86)
LIGHT = RGBColor(0x9b, 0xa4, 0xb2)
ORANGE = RGBColor(0xb8, 0x4a, 0x33)
CREAM_HEX = 'F8F4E9'
BORDER_HEX = 'E4DBC7'
WHITE = RGBColor(0xff, 0xff, 0xff)
WRAP = RGBColor(0xf2, 0xe8, 0xc9)

EA = 'Hiragino Sans GB'
LAT = 'Arial'

# ---------------------------------------------------------------- helpers
def set_font(run, size=10, bold=None, color=None, italic=None, latin=LAT, ea=EA):
    run.font.name = latin
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), latin)
    rFonts.set(qn('w:hAnsi'), latin)
    rFonts.set(qn('w:eastAsia'), ea)
    run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if italic is not None: run.font.italic = italic
    if color is not None: run.font.color.rgb = color

def char_spacing(run, twips):
    rPr = run._element.get_or_add_rPr()
    sp = OxmlElement('w:spacing'); sp.set(qn('w:val'), str(twips)); rPr.append(sp)

def border(paragraph_or_cell, edge, sz, color, space=0):
    pPr = paragraph_or_cell._p.get_or_add_pPr() if hasattr(paragraph_or_cell, '_p') else paragraph_or_cell._tc.get_or_add_tcPr()
    tag = 'w:pBdr' if hasattr(paragraph_or_cell, '_p') else 'w:tcBorders'
    el = pPr.find(qn(tag))
    if el is None:
        el = OxmlElement(tag); pPr.append(el)
    e = OxmlElement('w:' + edge)
    e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz))
    e.set(qn('w:space'), str(space)); e.set(qn('w:color'), color)
    el.append(e)

def bottom_rule(p, color=GOLD_HEX, sz=6):
    border(p, 'bottom', sz, color, 4)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:shd')): tcPr.remove(old)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def cell_borders(cell, color=BORDER_HEX, sz=4, edges=('top','left','bottom','right')):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:tcBorders')): tcPr.remove(old)
    tb = OxmlElement('w:tcBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement('w:' + edge)
        if edge in edges:
            e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), '0'); e.set(qn('w:color'), color)
        else:
            e.set(qn('w:val'), 'nil')
        tb.append(e)
    tcPr.append(tb)

def table_borders(table, color=BORDER_HEX, sz=4):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')): tblPr.remove(old)
    tb = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), '0'); e.set(qn('w:color'), color)
        tb.append(e)
    tblPr.append(tb)

def table_no_borders(table):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')): tblPr.remove(old)
    tb = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'nil'); tb.append(e)
    tblPr.append(tb)

def table_width(table, cm):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn('w:tblW')): tblPr.remove(old)
    w = OxmlElement('w:tblW'); w.set(qn('w:w'), str(int(cm / 2.54 * 1440))); w.set(qn('w:type'), 'dxa')
    tblPr.append(w)
    layout = OxmlElement('w:tblLayout'); layout.set(qn('w:type'), 'fixed'); tblPr.append(layout)

def table_cell_margins(table, top=60, bottom=60, left=100, right=100):
    tblPr = table._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for side, val in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement('w:' + side); e.set(qn('w:w'), str(val)); e.set(qn('w:type'), 'dxa'); mar.append(e)
    tblPr.append(mar)

def row_header(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); trPr.append(h)

def row_nosplit(row):
    trPr = row._tr.get_or_add_trPr()
    c = OxmlElement('w:cantSplit'); c.set(qn('w:val'), 'true'); trPr.append(c)

def cell_text(cell, text, size=8.6, bold=False, color=BODY, align=WD_ALIGN_PARAGRAPH.CENTER, italic=False):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.text = ''
    p.alignment = align
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text); set_font(r, size, bold, color, italic)

def add_picture(doc, path, width_cm, caption=None, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(4); pf.space_after = Pt(2); pf.line_spacing = 1.0
    pf.keep_with_next = bool(caption)
    p.add_run().add_picture(path, width=Cm(width_cm))
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(2); cp.paragraph_format.space_after = Pt(space_after)
        r = cp.add_run(caption); set_font(r, 8, False, GRAY, True)

def add_body(doc, text, size=10, after=5, before=0, color=BODY, line=1.42, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = line
    r = p.add_run(text); set_font(r, size, False, color)
    return p

def add_bullet(doc, text, size=9.6):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.55); pf.first_line_indent = Cm(-0.45)
    pf.space_before = Pt(0); pf.space_after = Pt(2.5); pf.line_spacing = 1.36
    r = p.add_run('·  '); set_font(r, size, True, GOLD_LIGHT)
    r2 = p.add_run(text); set_font(r2, size, False, BODY)
    return p

def add_kicker(doc, text, new_page=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.page_break_before = new_page
    pf.space_before = Pt(0); pf.space_after = Pt(3); pf.keep_with_next = True
    r = p.add_run(text); set_font(r, 8, True, GOLD); char_spacing(r, 50)

def add_h1(doc, cn, en):
    p = doc.add_paragraph(style='Heading 1')
    pf = p.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(0); pf.keep_with_next = True
    r = p.add_run(cn); set_font(r, 21, True, NAVY)
    enp = doc.add_paragraph()
    epf = enp.paragraph_format
    epf.space_before = Pt(3); epf.space_after = Pt(9); epf.keep_with_next = True
    r2 = enp.add_run(en); set_font(r2, 7.6, True, GRAY); char_spacing(r2, 80)
    bottom_rule(enp, GOLD_HEX, 6)

def add_h2(doc, text, before=10):
    p = doc.add_paragraph(style='Heading 2')
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(4); pf.keep_with_next = True
    r = p.add_run(text); set_font(r, 11.5, True, NAVY)

def add_note(doc, text, size=8.4, before=3, after=6):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = 1.3
    r = p.add_run(text); set_font(r, size, False, GRAY)
    return p

def page_break(doc):
    doc.add_page_break()

def add_page_field(paragraph):
    run = paragraph.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = ' PAGE '
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t'); t.text = '1'
    f3 = OxmlElement('w:fldChar'); f3.set(qn('w:fldCharType'), 'end')
    for el in (f1, it, f2, t, f3): run._r.append(el)
    set_font(run, 7.4, False, GRAY)

# ---------------------------------------------------------------- document setup
doc = Document()
doc.core_properties.title = '钢铁防线 3D · 坦克档案'
doc.core_properties.subject = 'Steel Bastion 3D Tank Dossier'
doc.core_properties.author = 'Steel Bastion 3D'
doc.core_properties.keywords = '坦克, Three.js, 游戏, 档案, 属性'

st = doc.styles['Normal']
st.font.name = LAT; st.font.size = Pt(10); st.font.color.rgb = BODY
st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), EA)
st.paragraph_format.line_spacing = 1.42
st.paragraph_format.space_after = Pt(5)

hs1 = doc.styles['Heading 1']
hs1.font.name = LAT; hs1.font.size = Pt(21); hs1.font.bold = True; hs1.font.color.rgb = NAVY
hs1.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), EA)
hs2 = doc.styles['Heading 2']
hs2.font.name = LAT; hs2.font.size = Pt(11.5); hs2.font.bold = True; hs2.font.color.rgb = NAVY
hs2.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), EA)

sec = doc.sections[0]
sec.page_width = Cm(21); sec.page_height = Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.95)
sec.top_margin = Cm(1.95); sec.bottom_margin = Cm(1.85)
sec.header_distance = Cm(0.9); sec.footer_distance = Cm(0.8)
sec.different_first_page_header_footer = True

# header / footer (not on cover)
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = hp.add_run('STEEL BASTION 3D   /   坦克档案'); set_font(r, 7.2, True, LIGHT); char_spacing(r, 30)
bottom_rule(hp, 'E4DBC7', 4)

try:
    doc.styles['Footer'].paragraph_format.tab_stops.clear_all()
except Exception:
    pass
fp = sec.footer.paragraphs[0]
fp.paragraph_format.tab_stops.clear_all()
fp.paragraph_format.tab_stops.add_tab_stop(Cm(17.05), WD_TAB_ALIGNMENT.RIGHT)
r = fp.add_run('STEEL BASTION 3D  ·  TANK DOSSIER'); set_font(r, 7, False, LIGHT)
r = fp.add_run('\t'); set_font(r, 7)
r = fp.add_run('PAGE  '); set_font(r, 7, False, LIGHT)
add_page_field(fp)

# ---------------------------------------------------------------- cover
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
p.paragraph_format.line_spacing = 1.0
p.add_run().add_picture(A('cover.png'), width=Cm(16.55))

# ---------------------------------------------------------------- 01 overview
add_kicker(doc, '01  /  OVERVIEW', True)
add_h1(doc, '车辆概览', 'THREE PLAYABLE TANKS  ·  TRAINING GROUNDS')
add_body(doc, '《钢铁防线 3D》当前提供三台可驾驶坦克，覆盖二战中型、二战重型与现代主战三个时代。三台车共用同一套驾驶、瞄准与开火逻辑，真正的差异在于装甲、机动、火力节奏，以及各自的一项战术技能。')
add_body(doc, '下面的数据直接读取自游戏配置表，渲染画面由游戏实时截取。除特别标注外，速度与弹速单位为 u/s（游戏内格/秒），转向单位为 rad/s。')

add_h2(doc, '数据对照', 8)
rows = [
    ('定位', '中型坦克 · 二战', '重型坦克 · 二战', '主战坦克 · 现代'),
    ('生命值', '100', '160', '120'),
    ('有效生命', '100', '200（-20% 受击）', '120 + 3 次免伤'),
    ('最大速度', '3.0 u/s', '1.8 u/s', '2.6 u/s'),
    ('车体转向', '2.8 rad/s', '1.6 rad/s', '2.4 rad/s'),
    ('装填时间', '0.50 s', '1.20 s', '0.40 s'),
    ('单发伤害', '25', '60', '30'),
    ('弹速', '9 u/s', '10 u/s', '12 u/s'),
    ('理论 DPS', '50', '50', '75'),
    ('技能', '烟幕掩护 · 20s', '穿甲弹 · 15s', '反应装甲 · 25s'),
    ('特性', '拾取补给后自修 5s', '受到伤害 -20%', '森林敌军不隐身'),
]
widths = [3.4, 4.55, 4.55, 4.55]
t = doc.add_table(rows=len(rows) + 1, cols=4)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
table_width(t, sum(widths)); table_borders(t, BORDER_HEX, 4); table_cell_margins(t, 58, 58, 110, 110)
for i, w in enumerate(widths):
    t.columns[i].width = Cm(w)
for row in t.rows:
    for i, w in enumerate(widths):
        row.cells[i].width = Cm(w)
    row_nosplit(row)
header = ('属性', 'M4 SHERMAN', 'TIGER I', 'T-90')
for i, h in enumerate(header):
    c = t.cell(0, i); shade(c, NAVY_HEX)
    cell_text(c, h, 8.0, True, WHITE, WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER)
row_header(t.rows[0])
for ri, data in enumerate(rows, start=1):
    for ci, val in enumerate(data):
        c = t.cell(ri, ci)
        if ri % 2 == 0: shade(c, 'FAF7EF')
        if ci == 0:
            cell_text(c, val, 8.5, True, NAVY, WD_ALIGN_PARAGRAPH.LEFT)
        else:
            cell_text(c, val, 8.5, False, BODY, WD_ALIGN_PARAGRAPH.CENTER)
add_note(doc, '有效生命 = 生命值 ÷ (1 - 减伤)；T-90 的反应装甲按“3 次伤害免疫”计。表中 DPS 为单发伤害 ÷ 装填时间，未计入穿甲弹与烟幕带来的输出窗口。', 8.2, 4, 4)
add_body(doc, '三台车形成了清晰的三角关系：Sherman 用机动和射速换取持续输出与烟幕窗口；Tiger I 用速度换装甲和单发爆发；T-90 则以最高 DPS、反应装甲和反伏击观瞄成为最稳的持续火力平台。', 9.4, 4)

add_h2(doc, '定位速览', 7)
cards = [
    ('M4 SHERMAN', '机动 · 均衡', '烟幕开道，持续游走', '6B7D4F'),
    ('TIGER I', '重装 · 攻坚', '正面推进，一炮破墙', '8A8D6D'),
    ('T-90', '速射 · 视野', '最高 DPS，反伏击', '4A5240'),
]
t = doc.add_table(rows=1, cols=3)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
table_width(t, 17.0); table_no_borders(t); table_cell_margins(t, 90, 90, 110, 110)
for i, w in enumerate((5.66, 5.66, 5.66)):
    t.columns[i].width = Cm(w)
for i, (name, role, line, color) in enumerate(cards):
    c = t.cell(0, i); c.width = Cm(5.66)
    shade(c, 'FAF7EF'); cell_borders(c, 'EFE6D3', 4)
    p = c.paragraphs[0]; p.text = ''
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(2)
    sq = p.add_run('■  '); set_font(sq, 8, True, RGBColor.from_string(color))
    r = p.add_run(name); set_font(r, 8.6, True, NAVY)
    p2 = c.add_paragraph(); p2.paragraph_format.space_before = Pt(0); p2.paragraph_format.space_after = Pt(2)
    r = p2.add_run(role); set_font(r, 7.8, True, GOLD)
    p3 = c.add_paragraph(); p3.paragraph_format.space_before = Pt(0); p3.paragraph_format.space_after = Pt(0)
    r = p3.add_run(line); set_font(r, 7.8, False, GRAY)

add_body(doc, '怎么选：要节奏与烟幕选 Sherman；要正面攻坚选 Tiger I；要射速与反伏击选 T-90。', 9.2, 0, before=4)

# ---------------------------------------------------------------- 02 stats
add_kicker(doc, '02  /  CORE STATS', True)
add_h1(doc, '核心属性', 'SIDE BY SIDE  ·  WHAT THE NUMBERS SAY')
add_picture(doc, A('chart.png'), 15.2, '图 1  ·  三台坦克的核心属性对比')
add_h2(doc, '三车合影', 7)
add_picture(doc, A('lineup.jpg'), 15.2, '图 2  ·  从左至右：M4 Sherman、Tiger I、T-90')

# ---------------------------------------------------------------- tank pages
def add_stat_strip(doc, items):
    t = doc.add_table(rows=2, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    table_width(t, 17.0); table_no_borders(t); table_cell_margins(t, 65, 65, 80, 80)
    widths = [5.66, 5.66, 5.66]
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    for ri in range(2):
        for ci in range(3):
            c = t.cell(ri, ci)
            c.width = Cm(widths[ci])
            shade(c, CREAM_HEX)
            cell_borders(c, 'EFE6D3', 6)
            cell_text(c, '', 8)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            val, label = items[ri * 3 + ci]
            r = p.add_run(val); set_font(r, 18, True, NAVY)
            p2 = c.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.paragraph_format.space_before = Pt(0); p2.paragraph_format.space_after = Pt(0)
            r2 = p2.add_run(label); set_font(r2, 7.2, False, GRAY)

def add_skill_box(doc, title, cd, desc, hint):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    table_width(t, 17.0); table_no_borders(t); table_cell_margins(t, 105, 105, 180, 160)
    c = t.cell(0, 0); c.width = Cm(17.0)
    shade(c, CREAM_HEX)
    cell_borders(c, GOLD_HEX, 4, edges=('top','bottom','right'))
    border(c, 'left', 18, GOLD_HEX, 0)
    p = c.paragraphs[0]; p.text = ''
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(3)
    r = p.add_run('技能 [E]   ·   ' + title); set_font(r, 11.2, True, NAVY)
    r2 = p.add_run('        冷却 ' + cd); set_font(r2, 8, False, GRAY)
    p2 = c.add_paragraph(); p2.paragraph_format.space_before = Pt(0); p2.paragraph_format.space_after = Pt(3)
    p2.paragraph_format.line_spacing = 1.35
    r = p2.add_run(desc); set_font(r, 9.2, False, BODY)
    p3 = c.add_paragraph(); p3.paragraph_format.space_before = Pt(0); p3.paragraph_format.space_after = Pt(0)
    p3.paragraph_format.line_spacing = 1.3
    r = p3.add_run(hint); set_font(r, 8.0, False, GRAY, True)

def add_pros_cons(doc, pros, cons):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    table_width(t, 17.0); table_no_borders(t); table_cell_margins(t, 60, 60, 110, 110)
    for i, w in enumerate((8.5, 8.5)):
        t.columns[i].width = Cm(w); t.cell(0, i).width = Cm(w)
        shade(t.cell(0, i), 'FAF7EF')
        cell_borders(t.cell(0, i), 'EFE6D3', 4)
    for i, (label, text, color) in enumerate((('优势', pros, GOLD), ('注意', cons, ORANGE))):
        c = t.cell(0, i)
        p = c.paragraphs[0]; p.text = ''
        p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(1)
        r = p.add_run(label); set_font(r, 7.6, True, color); char_spacing(r, 30)
        p2 = c.add_paragraph(); p2.paragraph_format.space_before = Pt(0); p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.line_spacing = 1.3
        r2 = p2.add_run(text); set_font(r2, 8.1, False, BODY)

def tank_page(kicker, name, en, cls, image, caption, stats, skill, role, pros, cons):
    add_kicker(doc, kicker, True)
    add_h1(doc, name, en)
    sub = doc.add_paragraph()
    spf = sub.paragraph_format
    spf.space_before = Pt(0); spf.space_after = Pt(4)
    r = sub.add_run(cls); set_font(r, 9.2, False, GRAY)
    add_picture(doc, A(image), 15.2, caption)
    add_stat_strip(doc, stats)
    add_skill_box(doc, skill[0], skill[1], skill[2], skill[3])
    add_h2(doc, '战术定位', 8)
    add_body(doc, role, 9.4, 3, line=1.36)
    add_pros_cons(doc, pros, cons)

tank_page(
    '03  /  VEHICLE 01', 'M4 Sherman', 'MEDIUM TANK  ·  WWII  ·  ALL-ROUNDER',
    '中型坦克 · 二战 · 均衡型', 'sherman.jpg', '图 3  ·  M4 Sherman 战场渲染',
    [('100', '生命值 HP'), ('3.0', '最大速度 u/s'), ('0.5s', '装填时间'),
     ('25', '单发伤害'), ('9', '弹速 u/s'), ('50', '理论 DPS')],
    ('烟幕掩护', '20 秒', '释放烟幕 5 秒，敌方 AI 无法锁定你。烟幕会同时遮断视线与射击，可用来强行穿越开阔地、掩护撤退或拆解狙击车的射界。', '提示：拾取任意补给后，Sherman 还会获得 5 秒自修（约 +10 HP）。'),
    '三车中最均衡的一台。3.0 u/s 的速度与 0.5 秒装填让它能持续游走、快速补位；25 点单发伤害偏低，面对 75 HP 的主战线车需要连续命中。烟幕既能掩护撤退，也能遮断狙击车的射界；每次拾取补给还会触发 5 秒自修（约 +10 HP），续航是三车中最好的。',
    '机动灵活，射速快，烟幕可攻可守，续航好', '单发伤害低，破墙效率一般，正面换血吃亏'
)

tank_page(
    '04  /  VEHICLE 02', 'Tiger I', 'HEAVY TANK  ·  WWII  ·  BREACHER',
    '重型坦克 · 二战 · 攻坚型', 'tiger.jpg', '图 4  ·  Tiger I 战场渲染',
    [('160', '生命值 HP'), ('1.8', '最大速度 u/s'), ('1.2s', '装填时间'),
     ('60', '单发伤害'), ('10', '弹速 u/s'), ('50', '理论 DPS')],
    ('穿甲弹', '15 秒', '下一炮击穿钢墙，并贯穿最多 2 个敌人。适合在僵持线上直接拆掉钢板掩体，或对准一条直线上的敌军制造双杀。', '提示：Tiger I 受到的所有伤害额外 -20%，有效生命约 200。'),
    '以装甲和单发伤害换机动的重装平台。160 HP 配合 20% 受击减免，实际可承受约 200 点伤害；60 点单发伤害两炮即可摧毁主战线车。1.2 秒装填意味着每一次开火都要有把握，穿甲弹则提供了击穿钢墙、贯穿最多 2 个敌人的破阵能力。',
    '有效生命最高，单发爆发强，穿甲弹可破墙/贯穿', '装填慢、转向慢，侧翼被绕时很难脱身'
)

tank_page(
    '05  /  VEHICLE 03', 'T-90', 'MAIN BATTLE TANK  ·  MODERN  ·  SUSTAINED FIRE',
    '主战坦克 · 现代 · 速射 / 视野型', 't90.jpg', '图 5  ·  T-90 战场渲染',
    [('120', '生命值 HP'), ('2.6', '最大速度 u/s'), ('0.4s', '装填时间'),
     ('30', '单发伤害'), ('12', '弹速 u/s'), ('75', '理论 DPS')],
    ('反应装甲', '25 秒', '吸收接下来 3 次伤害（30 秒内）。可以主动探头换血、吃掉狙击车的首发，或在冲锋时抵消一轮交叉火力。', '提示：T-90 的观瞄不受森林遮蔽影响，林中的敌军不会进入半透明隐身状态。'),
    '现代主战平台，用射速与观瞄换取容错。0.4 秒装填带来 75 点理论 DPS，是三车中最高的持续输出；120 HP 配合 3 层反应装甲，可以主动吃下 3 次伤害。它的观瞄优势让森林中的敌军不再进入隐身状态，反伏击能力最强。',
    'DPS 最高，反应装甲可主动换血，反伏击', '单发低，装甲层数有限，技能冷却较长'
)

# ---------------------------------------------------------------- 06 enemies
add_kicker(doc, '06  /  ENEMY FORCES', True)
add_h1(doc, '敌军单位', 'THREE HOSTILE CHASSIS  ·  WAVES 1—3')
add_body(doc, '敌军由三种车体轮换组成，每波最多同时存在 4 辆，三波共 10 辆。生命值以 25 点为一档，表中同时给出档位与原值。')
erows = [
    ('Scout 侦察车', '25（1）', '3.6', '1.5 s', '10', '5', '100', '0.8× 车体；未锁定时左右绕行搜索'),
    ('Line 主战线车', '75（3）', '1.6', '2.0 s', '20', '7', '200', '1.0× 车体；直线推进，持续压制'),
    ('Sniper 狙击车', '50（2）', '2.2', '2.5 s', '30', '9', '300', '0.9× 车体；长身管，0.55 s 瞄准，高精度'),
]
ew = [3.0, 1.75, 1.5, 1.75, 1.4, 1.35, 1.35, 4.9]
t = doc.add_table(rows=4, cols=8)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
table_width(t, sum(ew)); table_borders(t, BORDER_HEX, 4); table_cell_margins(t, 65, 65, 90, 90)
for i, w in enumerate(ew):
    t.columns[i].width = Cm(w)
for row in t.rows:
    for i, w in enumerate(ew):
        row.cells[i].width = Cm(w)
    row_nosplit(row)
headers = ('单位', '生命', '速度', '射击间隔', '伤害', '射程', '分值', '行为特征')
for i, h in enumerate(headers):
    c = t.cell(0, i); shade(c, NAVY_HEX)
    cell_text(c, h, 7.8, True, WHITE, WD_ALIGN_PARAGRAPH.LEFT if i in (0, 7) else WD_ALIGN_PARAGRAPH.CENTER)
row_header(t.rows[0])
for ri, data in enumerate(erows, start=1):
    for ci, val in enumerate(data):
        c = t.cell(ri, ci)
        if ri % 2 == 0: shade(c, 'FAF7EF')
        cell_text(c, val, 8.0, ci == 0, NAVY if ci == 0 else BODY,
                  WD_ALIGN_PARAGRAPH.LEFT if ci in (0, 7) else WD_ALIGN_PARAGRAPH.CENTER)
add_picture(doc, A('enemies.jpg'), 14.6, '图 6  ·  从左至右：Scout、Line、Sniper')
add_h2(doc, '行为逻辑', 7)
for b in (
    '敌军优先追击玩家；失去视线后改向鹰徽推进，形成“逼你回防”的压力。',
    '狙击车在 4.2 格内会后撤，在 4.2—9 格之间停车瞄准；0.55 秒瞄准后以极小误差射击。',
    '侦察车没有目标时会左右摆动搜索；被卡住 0.35 秒后尝试 90° 转向脱困。',
    '敌军进入森林会变为半透明；T-90 的观瞄不受影响。敌方炮弹速度为 7.5 u/s。',
):
    add_bullet(doc, b, 9.3)

# ---------------------------------------------------------------- 07 battlefield
add_kicker(doc, '07  /  BATTLEFIELD', True)
add_h1(doc, '战场与光影', 'PROCEDURAL PBR  ·  REALTIME LIGHTING')
add_body(doc, '当前版本对场景做了一次完整的渲染升级：天空、材质、阴影和粒子全部程序化生成，不依赖任何外部贴图。')
add_picture(doc, A('battle.jpg'), 15.0, '图 7  ·  战斗视角：追尾镜头、沙地、砖墙与鹰徽')
grid = doc.add_table(rows=1, cols=2)
grid.alignment = WD_TABLE_ALIGNMENT.CENTER
grid.autofit = False
table_width(grid, 17.0); table_no_borders(grid); table_cell_margins(grid, 0, 0, 40, 40)
for i, w in enumerate((8.2, 8.2)):
    grid.columns[i].width = Cm(w); grid.cell(0, i).width = Cm(w)
for i, (img, cap) in enumerate(((('topdown.jpg'), '图 8  ·  俯视视角'), (('smoke.jpg'), '图 9  ·  烟幕掩护'))):
    c = grid.cell(0, i)
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(A(img), width=Cm(7.7))
    cp = c.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(1); cp.paragraph_format.space_after = Pt(2)
    r = cp.add_run(cap); set_font(r, 7.6, False, GRAY, True)
add_h2(doc, '渲染要点', 7)
for b in (
    '天空与 IBL：程序化天空 + 太阳光晕生成 PMREM 环境贴图，金属与车体都有真实反射。',
    'PBR 材质：沙地、砖墙、钢板、迷彩与履带均带 albedo / normal / roughness 贴图。',
    '阴影：2048 PCF 软阴影 + 车底接触阴影；主光、补光、轮廓光与半球光四层布光。',
    '物理反馈：炮管后坐、履带滚动、扬尘、履带印、炮口硝烟与爆炸点光源。',
    '性能：1440×900 下约 60 FPS，材质与特效按战斗生命周期释放。',
):
    add_bullet(doc, b, 9.0)

# ---------------------------------------------------------------- 08 manual
add_kicker(doc, '08  /  FIELD MANUAL', True)
add_h1(doc, '操作与规则', 'CONTROLS  ·  WAVES  ·  POWER-UPS')
add_h2(doc, '操作', 2)
controls = [
    ('前进 / 后退', 'W / S'),
    ('车体转向', 'A / D'),
    ('炮塔瞄准', '鼠标'),
    ('开火', '空格 / 鼠标左键'),
    ('技能', 'E'),
    ('切换视角', 'C（追尾 / 俯视）'),
    ('返回车库', 'R'),
]
t = doc.add_table(rows=len(controls) + 1, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
wc = (5.0, 12.0)
table_width(t, sum(wc)); table_borders(t, BORDER_HEX, 4); table_cell_margins(t, 40, 40, 110, 110)
for i, w in enumerate(wc):
    t.columns[i].width = Cm(w)
for row in t.rows:
    for i, w in enumerate(wc): row.cells[i].width = Cm(w)
    row_nosplit(row)
for i, h in enumerate(('操作', '按键')):
    c = t.cell(0, i); shade(c, NAVY_HEX); cell_text(c, h, 8.2, True, WHITE, WD_ALIGN_PARAGRAPH.LEFT)
row_header(t.rows[0])
for ri, (a, b) in enumerate(controls, start=1):
    for ci, val in enumerate((a, b)):
        c = t.cell(ri, ci)
        if ri % 2 == 0: shade(c, 'FAF7EF')
        cell_text(c, val, 8.4, ci == 0, NAVY if ci == 0 else BODY, WD_ALIGN_PARAGRAPH.LEFT)

add_h2(doc, '规则', 6)
for b in (
    '守住底部鹰徽；任何炮弹——包括你自己的——命中鹰徽即告失败。',
    '三波敌军共 10 辆，波次在 0 / 15 / 30 秒触发，场上同时最多 4 辆。',
    '玩家 3 条命；阵亡 1.5 秒后重生，重生后 2.5 秒无敌。',
    '击毁得分：侦察车 100 / 主战线车 200 / 狙击车 300；拾取补给 +500；最高分写入浏览器 localStorage。',
):
    add_bullet(doc, b, 9.0)

add_h2(doc, '补给', 6)
powerups = [
    ('★  Star', '射速提升：装填时间 ×0.7（最低 0.22 s），并 +500 分'),
    ('Grenade', '立即摧毁场上所有敌军，+500 分'),
    ('Shovel', '鹰徽周围砖墙临时加固 15 秒，+500 分'),
    ('Clock', '全场敌军冻结 8 秒，+500 分'),
    ('Helmet', '能量护盾 10 秒，+500 分'),
    ('Tank', '备用坦克 +1'),
]
t = doc.add_table(rows=len(powerups) + 1, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
wp = (4.2, 12.8)
table_width(t, sum(wp)); table_borders(t, BORDER_HEX, 4); table_cell_margins(t, 40, 40, 110, 110)
for i, w in enumerate(wp):
    t.columns[i].width = Cm(w)
for row in t.rows:
    for i, w in enumerate(wp): row.cells[i].width = Cm(w)
    row_nosplit(row)
for i, h in enumerate(('补给', '效果')):
    c = t.cell(0, i); shade(c, NAVY_HEX); cell_text(c, h, 8.2, True, WHITE, WD_ALIGN_PARAGRAPH.LEFT)
row_header(t.rows[0])
for ri, (a, b) in enumerate(powerups, start=1):
    for ci, val in enumerate((a, b)):
        c = t.cell(ri, ci)
        if ri % 2 == 0: shade(c, 'FAF7EF')
        cell_text(c, val, 8.4, ci == 0, NAVY if ci == 0 else BODY, WD_ALIGN_PARAGRAPH.LEFT)
add_note(doc, '补给由被击毁的敌军以 22% 概率掉落，12 秒后消失。M4 Sherman 拾取任意补给后额外获得 5 秒自修。', 8.1, 4, 3)
add_note(doc, 'v1.0 · 2026.10  ·  场景光照与材质重制；三车技能修复；履带增加负重轮与履带板细节。', 7.6, 0, 2)
add_note(doc, '数据与素材：index.html · TANKS / ENEMY_TYPES / POWERUPS；Three.js r160（MIT 许可）。画面为游戏内实时渲染。', 7.4, 0, 0)

# ---------------------------------------------------------------- save
out = os.path.join(ROOT, 'docs', 'design', '钢铁防线3D_坦克档案.docx')
doc.save(out)
print('saved', out)
